from flask import Flask, request, render_template, send_file
from werkzeug.utils import secure_filename
import os
import tempfile
from docx import Document
from fpdf import FPDF
import pandas as pd
from PIL import Image

app = Flask(__name__)

PDF_FOLDER = './pdfs'
os.makedirs(PDF_FOLDER, exist_ok=True)
os.makedirs('./fonts', exist_ok=True)

def convert_docx_to_pdf(docx_path, pdf_path):
    doc = Document(docx_path)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font('DejaVu', '', 'fonts/DejaVuSans.ttf')
    pdf.set_font('DejaVu', size=10)

    for para in doc.paragraphs:
        if para.text.strip():
            pdf.multi_cell(190, 10, para.text)
            pdf.ln(5)

    pdf.output(pdf_path)

def convert_image_to_pdf(image_path, pdf_path):
    image = Image.open(image_path)
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font('DejaVu', '', 'fonts/DejaVuSans.ttf')
    pdf.set_font('DejaVu', size=10)
    pdf.image(image_path, x=10, y=10, w=190)
    pdf.output(pdf_path)

def convert_excel_to_pdf(excel_path, pdf_path):
    df = pd.read_excel(excel_path)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font('Arial', size=10)

    col_width = pdf.w / len(df.columns)
    row_height = 10

    for column in df.columns:
        pdf.cell(col_width, row_height, str(column), border=1, align='C')
    pdf.ln(row_height)

    for index, row in df.iterrows():
        for value in row:
            text = str(value)
            pdf.cell(col_width, row_height, text, border=1, align='C')
        pdf.ln(row_height)

    pdf.output(pdf_path)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        if 'file' not in request.files:
            return "No file part"
        file = request.files['file']
        if file.filename == '':
            return "No selected file"

        with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(file.filename)[1]) as tmp_file:
            filepath = tmp_file.name
            file.save(filepath)

        pdf_filename = secure_filename(file.filename.rsplit('.', 1)[0] + '.pdf')
        pdf_filepath = os.path.join(PDF_FOLDER, pdf_filename)
        file_extension = os.path.splitext(file.filename)[1].lower()

        try:
            if file_extension == '.docx':
                convert_docx_to_pdf(filepath, pdf_filepath)
            elif file_extension in ['.png', '.jpg', '.jpeg']:
                convert_image_to_pdf(filepath, pdf_filepath)
            elif file_extension in ['.xls', '.xlsx']:
                convert_excel_to_pdf(filepath, pdf_filepath)
            else:
                return "Unsupported file type"
        finally:
            if os.path.exists(filepath):
                os.remove(filepath)

        return send_file(pdf_filepath, as_attachment=True)
    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True)
